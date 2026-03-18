from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


def _parse_code_name(raw: object) -> tuple[str, str]:
    text = str(raw or "").strip()
    matched = re.match(r"^\((\d+)\)(.+)$", text)
    if matched:
        return matched.group(1), matched.group(2).strip()
    return "", text


def _normalize_mode(raw: object) -> str:
    text = str(raw or "").strip()
    if text == "非全":
        return "非全日制"
    return text


def _build_2024_summary(path: Path, top_schools: int, top_departments: int) -> tuple[dict, list[dict]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    header = [str(value).strip() if value is not None else "" for value in next(rows)]
    index = {name: pos for pos, name in enumerate(header)}

    school_counts: Counter[tuple[str, str]] = Counter()
    school_modes: defaultdict[tuple[str, str], Counter[str]] = defaultdict(Counter)
    department_counts: defaultdict[tuple[str, str], Counter[str]] = defaultdict(Counter)
    school_categories: dict[tuple[str, str], str] = {}
    total_rows = 0

    for row in rows:
        school_code, school_name = _parse_code_name(row[index["招生单位"]])
        if not school_name:
            continue
        department_code, department_name = _parse_code_name(row[index["院系所"]])
        school_key = (school_code, school_name)
        total_rows += 1
        school_counts[school_key] += 1
        if department_name:
            department_counts[school_key][department_name] += 1
        mode = _normalize_mode(row[index["学习方式"]])
        if mode:
            school_modes[school_key][mode] += 1
        school_categories.setdefault(school_key, "2024调剂余额表")

    summary_rows: list[dict] = []
    target_rows: list[dict] = []
    for rank, (school_key, count) in enumerate(school_counts.most_common(top_schools), start=1):
        school_code, school_name = school_key
        departments = department_counts[school_key].most_common(top_departments)
        summary_rows.append(
            {
                "rank": rank,
                "school_code": school_code,
                "school_name": school_name,
                "school_category": school_categories.get(school_key, ""),
                "adjustment_count": count,
                "study_modes": dict(school_modes[school_key].most_common()),
                "top_departments": [{"department_name": name, "count": dept_count} for name, dept_count in departments],
            }
        )
        if departments:
            for department_name, department_count in departments:
                target_rows.append(
                    {
                        "school_code": school_code,
                        "school_name": school_name,
                        "department_name": department_name,
                        "notes": f"2024调剂余额表 Top学校#{rank} / 学校样本{count} / 学院样本{department_count}",
                        "source": "adjustment_opportunity_2024",
                        "school_category": school_categories.get(school_key, ""),
                        "adjustment_count": count,
                    }
                )
        else:
            target_rows.append(
                {
                    "school_code": school_code,
                    "school_name": school_name,
                    "department_name": "未区分院系",
                    "notes": f"2024调剂余额表 Top学校#{rank} / 学校样本{count}",
                    "source": "adjustment_opportunity_2024",
                    "school_category": school_categories.get(school_key, ""),
                    "adjustment_count": count,
                }
            )

    summary = {
        "source_file": str(path),
        "sheet_name": worksheet.title,
        "total_rows": total_rows,
        "unique_schools": len(school_counts),
        "study_mode_counts": dict(sum((school_modes[key] for key in school_modes), Counter()).most_common()),
        "top_schools": summary_rows,
    }
    return summary, target_rows


def _build_2025_snapshot_summary(path: Path, top_schools: int) -> tuple[dict, list[dict]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    header = [str(value).strip() if value is not None else "" for value in next(rows)]
    index = {name: pos for pos, name in enumerate(header)}

    school_counts: Counter[str] = Counter()
    validation_states: Counter[str] = Counter()
    school_urls: defaultdict[str, Counter[str]] = defaultdict(Counter)
    major_counts: defaultdict[str, Counter[str]] = defaultdict(Counter)
    total_rows = 0

    for row in rows:
        if not any(row):
            continue
        school_name = str(row[index["学校"]] or "").strip()
        if not school_name:
            continue
        total_rows += 1
        school_counts[school_name] += 1
        state = str(row[index["验证状态"]] or "").strip()
        if state:
            validation_states[state] += 1
        url = str(row[index["原始网址"]] or "").strip()
        if url:
            school_urls[school_name][url] += 1
        major_label = f"{str(row[index['专业代码']] or '').strip()} {str(row[index['专业名称']] or '').strip()}".strip()
        if major_label:
            major_counts[school_name][major_label] += 1

    summary_rows: list[dict] = []
    target_rows: list[dict] = []
    for rank, (school_name, count) in enumerate(school_counts.most_common(top_schools), start=1):
        top_url = school_urls[school_name].most_common(1)
        top_majors = major_counts[school_name].most_common(3)
        summary_rows.append(
            {
                "rank": rank,
                "school_name": school_name,
                "school_category": "2025调剂信息快照",
                "adjustment_count": count,
                "top_departments": [],
                "top_majors": [{"major": name, "count": major_count} for name, major_count in top_majors],
                "top_url": top_url[0][0] if top_url else None,
            }
        )
        target_rows.append(
            {
                "school_name": school_name,
                "department_name": "未区分院系",
                "notes": f"2025调剂信息快照 Top学校#{rank} / 学校样本{count}",
                "source": "adjustment_opportunity_2025_snapshot",
                "school_category": "2025调剂信息快照",
                "adjustment_count": count,
                "top_url": top_url[0][0] if top_url else None,
            }
        )

    summary = {
        "source_file": str(path),
        "sheet_name": worksheet.title,
        "total_rows": total_rows,
        "unique_schools": len(school_counts),
        "validation_state_counts": dict(validation_states.most_common()),
        "top_schools": summary_rows,
    }
    return summary, target_rows


def build_outputs(
    *,
    adjustment_2024_path: Path,
    adjustment_2025_snapshot_path: Path,
    summary_2024_output: Path,
    summary_2025_output: Path,
    targets_output: Path,
    top_2024_schools: int,
    top_2024_departments: int,
    top_2025_schools: int,
) -> dict:
    summary_2024, targets_2024 = _build_2024_summary(adjustment_2024_path, top_2024_schools, top_2024_departments)
    summary_2025, targets_2025 = _build_2025_snapshot_summary(adjustment_2025_snapshot_path, top_2025_schools)

    deduped_targets: list[dict] = []
    seen_keys: set[tuple[str, str]] = set()
    for row in [*targets_2024, *targets_2025]:
        key = (str(row.get("school_name") or "").strip(), str(row.get("department_name") or "").strip())
        if not key[0] or key in seen_keys:
            continue
        seen_keys.add(key)
        deduped_targets.append(row)

    summary_2024_output.parent.mkdir(parents=True, exist_ok=True)
    summary_2025_output.parent.mkdir(parents=True, exist_ok=True)
    targets_output.parent.mkdir(parents=True, exist_ok=True)
    summary_2024_output.write_text(json.dumps(summary_2024, ensure_ascii=False, indent=2), encoding="utf-8")
    summary_2025_output.write_text(json.dumps(summary_2025, ensure_ascii=False, indent=2), encoding="utf-8")
    targets_output.write_text(json.dumps(deduped_targets, ensure_ascii=False, indent=2), encoding="utf-8")

    return {
        "summary_2024_output": str(summary_2024_output),
        "summary_2025_output": str(summary_2025_output),
        "targets_output": str(targets_output),
        "rows_2024": summary_2024["total_rows"],
        "rows_2025": summary_2025["total_rows"],
        "target_rows": len(deduped_targets),
        "target_unique_schools": len({row["school_name"] for row in deduped_targets}),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Build supplemental adjustment seed outputs from 2024/2025 workbooks.")
    parser.add_argument("--adjustment-2024", required=True, help="Path to 2024 adjustment workbook")
    parser.add_argument("--adjustment-2025-snapshot", required=True, help="Path to 2025 snapshot workbook")
    parser.add_argument(
        "--summary-2024-output",
        default="docs/data/adjustment_opportunity_2024_summary.json",
        help="Path to write the 2024 summary JSON",
    )
    parser.add_argument(
        "--summary-2025-output",
        default="docs/data/adjustment_opportunity_2025_snapshot_summary.json",
        help="Path to write the 2025 snapshot summary JSON",
    )
    parser.add_argument(
        "--targets-output",
        default="docs/data/adjustment_supplemental_priority_targets_2024_2025.json",
        help="Path to write the merged school target JSON",
    )
    parser.add_argument("--top-2024-schools", type=int, default=30)
    parser.add_argument("--top-2024-departments", type=int, default=2)
    parser.add_argument("--top-2025-schools", type=int, default=30)
    args = parser.parse_args()

    result = build_outputs(
        adjustment_2024_path=Path(args.adjustment_2024).expanduser().resolve(),
        adjustment_2025_snapshot_path=Path(args.adjustment_2025_snapshot).expanduser().resolve(),
        summary_2024_output=Path(args.summary_2024_output).resolve(),
        summary_2025_output=Path(args.summary_2025_output).resolve(),
        targets_output=Path(args.targets_output).resolve(),
        top_2024_schools=args.top_2024_schools,
        top_2024_departments=args.top_2024_departments,
        top_2025_schools=args.top_2025_schools,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
