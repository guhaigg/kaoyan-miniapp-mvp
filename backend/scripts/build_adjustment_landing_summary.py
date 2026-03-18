from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


def _parse_code_name(raw: object) -> tuple[str, str]:
    text = str(raw or "").strip()
    matched = re.match(r"^\((\d+)\)(.+)$", text)
    if matched:
        return matched.group(1), matched.group(2).strip()
    return "", text


def _normalize_mode(raw: object) -> str:
    text = str(raw or "").strip()
    if text == "非全":
        return "非全日制"
    return text


def build_summary(*, input_path: Path, output_path: Path, top_schools: int = 20) -> dict:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    header = [str(value).strip() if value is not None else "" for value in next(rows)]
    index = {name: pos for pos, name in enumerate(header)}

    school_counts: Counter[tuple[str, str]] = Counter()
    department_counts: defaultdict[tuple[str, str], Counter[str]] = defaultdict(Counter)
    first_choice_counts: Counter[str] = Counter()
    category_counts: Counter[str] = Counter()
    mode_counts: Counter[str] = Counter()
    score_bucket_counts: Counter[str] = Counter()
    total_rows = 0

    for row in rows:
        school_code, school_name = _parse_code_name(row[index["学校"]])
        if not school_name:
            continue
        total_rows += 1
        school_key = (school_code, school_name)
        school_counts[school_key] += 1

        department_name = str(row[index["所属学院"]] or "").strip() or "未区分院系"
        department_counts[school_key][department_name] += 1

        _, first_choice = _parse_code_name(row[index["一志愿报考院校"]])
        if first_choice:
            first_choice_counts[first_choice] += 1

        category = str(row[index["院校类别"]] or "").strip()
        if category:
            category_counts[category] += 1

        mode = _normalize_mode(row[index["学习形式"]])
        if mode:
            mode_counts[mode] += 1

        try:
            score = int(row[index["初试总分"]])
            bucket_floor = score // 25 * 25
            score_bucket_counts[f"{bucket_floor}-{bucket_floor + 24}"] += 1
        except Exception:
            pass

    top_school_rows: list[dict] = []
    for rank, (school_key, count) in enumerate(school_counts.most_common(top_schools), start=1):
        school_code, school_name = school_key
        departments = department_counts[school_key].most_common(3)
        top_school_rows.append(
            {
                "rank": rank,
                "school_code": school_code,
                "school_name": school_name,
                "school_category": "2025调剂上岸名单",
                "adjustment_count": count,
                "top_departments": [{"department_name": name, "count": dept_count} for name, dept_count in departments],
            }
        )

    summary = {
        "source_file": str(input_path),
        "sheet_name": worksheet.title,
        "total_rows": total_rows,
        "unique_schools": len(school_counts),
        "category_counts": dict(category_counts.most_common()),
        "study_mode_counts": dict(mode_counts.most_common()),
        "score_bucket_counts": dict(score_bucket_counts.most_common()),
        "top_first_choices": [{"school_name": name, "count": count} for name, count in first_choice_counts.most_common(10)],
        "top_schools": top_school_rows,
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description="Build adjustment landing summary from workbook.")
    parser.add_argument("--input", required=True, help="Path to the landing workbook")
    parser.add_argument(
        "--output",
        default="docs/data/adjustment_landing_2025_summary.json",
        help="Path to write the landing summary JSON",
    )
    parser.add_argument("--top-schools", type=int, default=20)
    args = parser.parse_args()

    summary = build_summary(
        input_path=Path(args.input).expanduser().resolve(),
        output_path=Path(args.output).resolve(),
        top_schools=args.top_schools,
    )
    print(
        json.dumps(
            {
                "total_rows": summary["total_rows"],
                "unique_schools": summary["unique_schools"],
                "output": str(Path(args.output).resolve()),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
