from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


def _parse_school(raw: object) -> tuple[str, str]:
    text = str(raw or "").strip()
    matched = re.match(r"^\((\d+)\)(.+)$", text)
    if matched:
        return matched.group(1), matched.group(2).strip()
    return "", text


def _parse_region(raw: object) -> tuple[str, str]:
    text = str(raw or "").strip()
    matched = re.match(r"^\((\d+)\)(.+)$", text)
    if matched:
        return matched.group(1), matched.group(2).strip()
    return "", text


def _normalize_mode(raw: object) -> str:
    text = str(raw or "").strip()
    if text == "非全":
        return "非全日制"
    return text


def build_outputs(
    *,
    input_path: Path,
    summary_output: Path,
    targets_output: Path,
    top_schools: int,
    top_departments: int,
) -> tuple[dict, list[dict]]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)

    next(rows, None)  # title row
    header = [str(value).strip() if value is not None else "" for value in next(rows)]
    index = {name: pos for pos, name in enumerate(header)}

    school_counts: Counter[tuple[str, str]] = Counter()
    school_categories: dict[tuple[str, str], str] = {}
    school_regions: dict[tuple[str, str], tuple[str, str]] = {}
    school_years: defaultdict[tuple[str, str], set[str]] = defaultdict(set)
    school_modes: defaultdict[tuple[str, str], Counter[str]] = defaultdict(Counter)
    department_counts: defaultdict[tuple[str, str], Counter[str]] = defaultdict(Counter)
    major_counts: defaultdict[tuple[str, str], Counter[str]] = defaultdict(Counter)
    year_counts: Counter[str] = Counter()
    mode_counts: Counter[str] = Counter()
    total_rows = 0

    for row in rows:
        school_code, school_name = _parse_school(row[index["学校"]])
        if not school_name:
            continue
        school_key = (school_code, school_name)
        region_code, region_name = _parse_region(row[index["地区"]])
        category = str(row[index["院校类别"]] or "").strip()
        department_name = str(row[index["所属学院"]] or "").strip() or "未区分院系"
        major_code = str(row[index["专业代码"]] or "").strip()
        major_name = str(row[index["专业名称"]] or "").strip()
        study_mode = _normalize_mode(row[index["学习形式"]])
        year = str(row[index["年份"]] or "").strip()

        total_rows += 1
        school_counts[school_key] += 1
        school_categories.setdefault(school_key, category)
        school_regions.setdefault(school_key, (region_code, region_name))
        if year:
            school_years[school_key].add(year)
            year_counts[year] += 1
        if study_mode:
            school_modes[school_key][study_mode] += 1
            mode_counts[study_mode] += 1
        if department_name:
            department_counts[school_key][department_name] += 1
        if major_code or major_name:
            major_label = f"{major_code} {major_name}".strip()
            major_counts[school_key][major_label] += 1

    top_school_rows = []
    target_rows = []
    for rank, (school_key, count) in enumerate(school_counts.most_common(top_schools), start=1):
        school_code, school_name = school_key
        region_code, region_name = school_regions.get(school_key, ("", ""))
        category = school_categories.get(school_key, "")
        departments = department_counts[school_key].most_common(top_departments)
        majors = major_counts[school_key].most_common(3)
        modes = dict(school_modes[school_key].most_common())
        years = sorted(school_years[school_key])

        top_school_rows.append(
            {
                "rank": rank,
                "school_code": school_code,
                "school_name": school_name,
                "region_code": region_code,
                "region_name": region_name,
                "school_category": category,
                "adjustment_count": count,
                "years": years,
                "study_modes": modes,
                "top_departments": [{"department_name": name, "count": dept_count} for name, dept_count in departments],
                "top_majors": [{"major": name, "count": major_count} for name, major_count in majors],
            }
        )

        if departments:
            for department_name, department_count in departments:
                target_rows.append(
                    {
                        "school_code": school_code,
                        "school_name": school_name,
                        "department_name": department_name,
                        "notes": f"调剂统计23-25 Top学校#{rank} / 学校样本{count} / 学院样本{department_count}",
                        "source": "adjustment_stats_2023_2025",
                        "region_name": region_name,
                        "school_category": category,
                        "adjustment_count": count,
                        "years": years,
                    }
                )
        else:
            target_rows.append(
                {
                    "school_code": school_code,
                    "school_name": school_name,
                    "department_name": "未区分院系",
                    "notes": f"调剂统计23-25 Top学校#{rank} / 学校样本{count}",
                    "source": "adjustment_stats_2023_2025",
                    "region_name": region_name,
                    "school_category": category,
                    "adjustment_count": count,
                    "years": years,
                }
            )

    summary = {
        "source_file": str(input_path),
        "sheet_name": worksheet.title,
        "total_rows": total_rows,
        "year_counts": dict(year_counts.most_common()),
        "study_mode_counts": dict(mode_counts.most_common()),
        "top_schools": top_school_rows,
    }

    summary_output.parent.mkdir(parents=True, exist_ok=True)
    targets_output.parent.mkdir(parents=True, exist_ok=True)
    summary_output.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    targets_output.write_text(json.dumps(target_rows, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary, target_rows


def main() -> None:
    parser = argparse.ArgumentParser(description="Build priority school targets from adjustment statistics workbook.")
    parser.add_argument("--input", required=True, help="Path to the source .xlsx workbook")
    parser.add_argument(
        "--summary-output",
        default="docs/data/adjustment_stats_2023_2025_summary.json",
        help="Path to write the aggregated summary JSON",
    )
    parser.add_argument(
        "--targets-output",
        default="docs/data/adjustment_priority_school_targets_2023_2025.json",
        help="Path to write the import-ready school target JSON",
    )
    parser.add_argument("--top-schools", type=int, default=60)
    parser.add_argument("--top-departments", type=int, default=2)
    args = parser.parse_args()

    summary, targets = build_outputs(
        input_path=Path(args.input).expanduser().resolve(),
        summary_output=Path(args.summary_output).resolve(),
        targets_output=Path(args.targets_output).resolve(),
        top_schools=args.top_schools,
        top_departments=args.top_departments,
    )
    print(
        json.dumps(
            {
                "total_rows": summary["total_rows"],
                "top_school_count": len(summary["top_schools"]),
                "target_rows": len(targets),
                "summary_output": str(Path(args.summary_output).resolve()),
                "targets_output": str(Path(args.targets_output).resolve()),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
