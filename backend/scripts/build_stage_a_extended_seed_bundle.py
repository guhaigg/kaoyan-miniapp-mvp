from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import xlrd
from openpyxl import load_workbook


def _parse_code_name(raw: object) -> tuple[str, str]:
    text = str(raw or "").strip()
    matched = re.match(r"^\((\d+)\)(.+)$", text)
    if matched:
        return matched.group(1), matched.group(2).strip()
    return "", text


def _parse_region(raw: object) -> str:
    text = str(raw or "").strip()
    matched = re.match(r"^\((\d+)\)(.+)$", text)
    if matched:
        return matched.group(2).strip()
    return text


def _normalize_mode(raw: object) -> str:
    text = str(raw or "").strip()
    if text in {"1", "全", "全日制"}:
        return "全日制"
    if text in {"2", "非全", "非全日制"}:
        return "非全日制"
    return text


def _contains_url(raw: object) -> bool:
    return "http://" in str(raw or "") or "https://" in str(raw or "")


def _load_first_sheet_xlsx(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    header = [str(value).strip() if value is not None else "" for value in next(rows)]
    return worksheet.title, header, rows


def _build_2024_snapshot_summary(path: Path, top_schools: int = 20, top_departments: int = 2) -> tuple[dict, list[dict]]:
    sheet_name, header, rows = _load_first_sheet_xlsx(path)
    index = {name: pos for pos, name in enumerate(header)}

    school_counts: Counter[tuple[str, str]] = Counter()
    department_counts: defaultdict[tuple[str, str], Counter[str]] = defaultdict(Counter)
    mode_counts: Counter[str] = Counter()
    url_rows = 0
    total_rows = 0

    for row in rows:
        school_code = str(row[index["学校代码"]] or "").strip()
        school_name = str(row[index["大学"]] or "").strip()
        if not school_name:
            continue
        department_name = str(row[index["学院"]] or "").strip() or "未区分院系"
        school_key = (school_code, school_name)
        total_rows += 1
        school_counts[school_key] += 1
        department_counts[school_key][department_name] += 1
        mode = _normalize_mode(row[index["学习形式"]])
        if mode:
            mode_counts[mode] += 1
        if _contains_url(row[index["要求"]]):
            url_rows += 1

    top_school_rows: list[dict] = []
    targets: list[dict] = []
    for rank, (school_key, count) in enumerate(school_counts.most_common(top_schools), start=1):
        school_code, school_name = school_key
        top_departments_rows = department_counts[school_key].most_common(top_departments)
        top_school_rows.append(
            {
                "rank": rank,
                "school_code": school_code,
                "school_name": school_name,
                "school_category": "2024调剂快照",
                "adjustment_count": count,
                "top_departments": [{"department_name": name, "count": dept_count} for name, dept_count in top_departments_rows],
            }
        )
        for department_name, dept_count in top_departments_rows:
            targets.append(
                {
                    "school_code": school_code,
                    "school_name": school_name,
                    "department_name": department_name,
                    "school_category": "2024调剂快照",
                    "notes": f"2024调剂快照 Top学校#{rank} / 学校样本{count} / 学院样本{dept_count}",
                    "source": "adjustment_snapshot_2024_0414",
                    "adjustment_count": count,
                }
            )

    summary = {
        "source_file": str(path),
        "sheet_name": sheet_name,
        "total_rows": total_rows,
        "unique_schools": len(school_counts),
        "study_mode_counts": dict(mode_counts.most_common()),
        "url_requirement_rows": url_rows,
        "top_schools": top_school_rows,
    }
    return summary, targets


def _build_2026_catalog_summary(path: Path, top_schools: int = 20) -> tuple[dict, list[dict]]:
    sheet_name, header, rows = _load_first_sheet_xlsx(path)
    index = {name: pos for pos, name in enumerate(header)}

    school_counts: Counter[str] = Counter()
    province_counts: Counter[str] = Counter()
    major_counts: defaultdict[str, Counter[str]] = defaultdict(Counter)
    total_slots = 0
    total_rows = 0

    for row in rows:
        school_name = str(row[index["院校名称"]] or "").strip()
        if not school_name:
            continue
        total_rows += 1
        school_counts[school_name] += 1
        province = str(row[index["省份"]] or "").strip()
        if province:
            province_counts[province] += 1
        major_name = f"{str(row[index['专业代码']] or '').strip()} {str(row[index['专业名称']] or '').strip()}".strip()
        if major_name:
            major_counts[school_name][major_name] += 1
        try:
            total_slots += int(row[index["名额"]] or 0)
        except Exception:
            pass

    top_school_rows: list[dict] = []
    targets: list[dict] = []
    for rank, (school_name, count) in enumerate(school_counts.most_common(top_schools), start=1):
        top_school_rows.append(
            {
                "rank": rank,
                "school_name": school_name,
                "school_category": "2026招生专业目录",
                "adjustment_count": count,
                "top_departments": [],
                "top_majors": [{"major": name, "count": major_count} for name, major_count in major_counts[school_name].most_common(3)],
            }
        )
        targets.append(
            {
                "school_name": school_name,
                "department_name": "未区分院系",
                "school_category": "2026招生专业目录",
                "notes": f"2026招生专业目录 Top学校#{rank} / 专业样本{count}",
                "source": "program_catalog_2026",
                "adjustment_count": count,
            }
        )

    summary = {
        "source_file": str(path),
        "sheet_name": sheet_name,
        "total_rows": total_rows,
        "unique_schools": len(school_counts),
        "province_counts": dict(province_counts.most_common()),
        "total_slots": total_slots,
        "top_schools": top_school_rows,
    }
    return summary, targets


def _build_2025_full_stats_summary(path: Path, top_schools: int = 20) -> tuple[dict, list[dict]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["调剂统计（含平均分）"]
    rows = worksheet.iter_rows(values_only=True)
    header = [str(value).strip() if value is not None else "" for value in next(rows)]
    index = {name: pos for pos, name in enumerate(header)}

    school_counts: Counter[tuple[str, str]] = Counter()
    mode_counts: Counter[str] = Counter()
    major_counts: defaultdict[tuple[str, str], Counter[str]] = defaultdict(Counter)
    total_people = 0
    min_scores: list[int] = []
    avg_scores: list[int] = []
    total_rows = 0

    for row in rows:
        school_code, school_name = _parse_code_name(row[index["推荐院校"]])
        if not school_name:
            continue
        total_rows += 1
        school_key = (school_code, school_name)
        school_counts[school_key] += 1
        mode = _normalize_mode(row[index["学习形式"]])
        if mode:
            mode_counts[mode] += 1
        major_name = str(row[index["专业"]] or "").strip()
        if major_name:
            major_counts[school_key][major_name] += 1
        try:
            total_people += int(row[index["25调剂人数"]] or 0)
        except Exception:
            pass
        for target_list, field_name in ((min_scores, "25调剂录取最低分"), (avg_scores, "25调剂平均分")):
            try:
                target_list.append(int(row[index[field_name]] or 0))
            except Exception:
                pass

    top_school_rows: list[dict] = []
    targets: list[dict] = []
    for rank, (school_key, count) in enumerate(school_counts.most_common(top_schools), start=1):
        school_code, school_name = school_key
        top_school_rows.append(
            {
                "rank": rank,
                "school_code": school_code,
                "school_name": school_name,
                "school_category": "2025调剂统计完整版",
                "adjustment_count": count,
                "top_departments": [],
                "top_majors": [{"major": name, "count": major_count} for name, major_count in major_counts[school_key].most_common(3)],
            }
        )
        targets.append(
            {
                "school_code": school_code,
                "school_name": school_name,
                "department_name": "未区分院系",
                "school_category": "2025调剂统计完整版",
                "notes": f"2025调剂统计完整版 Top学校#{rank} / 学校样本{count}",
                "source": "adjustment_stats_2025_full",
                "adjustment_count": count,
            }
        )

    summary = {
        "source_file": str(path),
        "sheet_name": worksheet.title,
        "total_rows": total_rows,
        "unique_schools": len(school_counts),
        "study_mode_counts": dict(mode_counts.most_common()),
        "total_people": total_people,
        "min_score_floor": min(min_scores) if min_scores else None,
        "avg_score_mean": round(sum(avg_scores) / len(avg_scores), 2) if avg_scores else None,
        "top_schools": top_school_rows,
    }
    return summary, targets


def _build_2025_announcement_summary(path: Path, top_schools: int = 20) -> tuple[dict, list[dict]]:
    workbook = xlrd.open_workbook(str(path))
    sheet = workbook.sheet_by_index(0)
    header = [str(value).strip() for value in sheet.row_values(0)]
    index = {name: pos for pos, name in enumerate(header)}

    school_counts: Counter[str] = Counter()
    verification_counts: Counter[str] = Counter()
    school_urls: defaultdict[str, Counter[str]] = defaultdict(Counter)
    total_rows = 0

    for row_idx in range(1, sheet.nrows):
        row = sheet.row_values(row_idx)
        school_name = str(row[index["学校"]] or "").strip()
        if not school_name:
            continue
        total_rows += 1
        school_counts[school_name] += 1
        verification = str(row[index["验证状态"]] or "").strip()
        if verification:
            verification_counts[verification] += 1
        url = str(row[index["原始网址"]] or "").strip()
        if url:
            school_urls[school_name][url] += 1

    top_school_rows: list[dict] = []
    targets: list[dict] = []
    for rank, (school_name, count) in enumerate(school_counts.most_common(top_schools), start=1):
        top_url = school_urls[school_name].most_common(1)
        top_school_rows.append(
            {
                "rank": rank,
                "school_name": school_name,
                "school_category": "2025调剂信息公告",
                "adjustment_count": count,
                "top_departments": [],
                "top_url": top_url[0][0] if top_url else None,
            }
        )
        targets.append(
            {
                "school_name": school_name,
                "department_name": "未区分院系",
                "school_category": "2025调剂信息公告",
                "notes": f"2025调剂信息公告 Top学校#{rank} / 公告样本{count}",
                "source": "adjustment_announcement_2025",
                "adjustment_count": count,
                "top_url": top_url[0][0] if top_url else None,
            }
        )

    summary = {
        "source_file": str(path),
        "sheet_name": workbook.sheet_names()[0],
        "total_rows": total_rows,
        "unique_schools": len(school_counts),
        "verification_counts": dict(verification_counts.most_common()),
        "top_schools": top_school_rows,
    }
    return summary, targets


def _build_2024_landing_summary(path: Path, top_schools: int = 20, top_departments: int = 2) -> tuple[dict, list[dict]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["总表"]
    rows = worksheet.iter_rows(values_only=True)
    header = [str(value).strip() if value is not None else "" for value in next(rows)]
    index = {name: pos for pos, name in enumerate(header)}

    school_counts: Counter[tuple[str, str]] = Counter()
    department_counts: defaultdict[tuple[str, str], Counter[str]] = defaultdict(Counter)
    category_counts: Counter[str] = Counter()
    region_counts: Counter[str] = Counter()
    total_rows = 0

    for row in rows:
        school_code, school_name = _parse_code_name(row[index["调剂学校"]])
        if not school_name:
            continue
        total_rows += 1
        school_key = (school_code, school_name)
        school_counts[school_key] += 1
        department_name = str(row[index["所属学院"]] or "").strip() or "未区分院系"
        department_counts[school_key][department_name] += 1
        category = str(row[index["院校类别"]] or "").strip()
        if category:
            category_counts[category] += 1
        region = _parse_region(row[index["地区"]])
        if region:
            region_counts[region] += 1

    top_school_rows: list[dict] = []
    targets: list[dict] = []
    for rank, (school_key, count) in enumerate(school_counts.most_common(top_schools), start=1):
        school_code, school_name = school_key
        top_departments_rows = department_counts[school_key].most_common(top_departments)
        top_school_rows.append(
            {
                "rank": rank,
                "school_code": school_code,
                "school_name": school_name,
                "school_category": "2024调剂上岸画像",
                "adjustment_count": count,
                "top_departments": [{"department_name": name, "count": dept_count} for name, dept_count in top_departments_rows],
            }
        )
        for department_name, dept_count in top_departments_rows:
            targets.append(
                {
                    "school_code": school_code,
                    "school_name": school_name,
                    "department_name": department_name,
                    "school_category": "2024调剂上岸画像",
                    "notes": f"2024调剂上岸画像 Top学校#{rank} / 学校样本{count} / 学院样本{dept_count}",
                    "source": "adjustment_landing_2024",
                    "adjustment_count": count,
                }
            )

    summary = {
        "source_file": str(path),
        "sheet_name": worksheet.title,
        "total_rows": total_rows,
        "unique_schools": len(school_counts),
        "category_counts": dict(category_counts.most_common()),
        "region_counts": dict(region_counts.most_common(10)),
        "top_schools": top_school_rows,
    }
    return summary, targets


def _dedupe_targets(rows: list[dict]) -> list[dict]:
    deduped: list[dict] = []
    seen_keys: set[tuple[str, str]] = set()
    for row in rows:
        school_name = str(row.get("school_name") or "").strip()
        department_name = str(row.get("department_name") or "").strip() or "未区分院系"
        if not school_name:
            continue
        key = (school_name, department_name)
        if key in seen_keys:
            continue
        seen_keys.add(key)
        deduped.append(row)
    return deduped


def build_outputs(
    *,
    adjustment_snapshot_2024_path: Path,
    program_catalog_2026_path: Path,
    adjustment_stats_2025_full_path: Path,
    adjustment_announcement_2025_path: Path,
    adjustment_landing_2024_path: Path,
    snapshot_2024_output: Path,
    program_catalog_output: Path,
    stats_2025_full_output: Path,
    announcement_2025_output: Path,
    landing_2024_output: Path,
    targets_output: Path,
) -> dict:
    snapshot_2024_summary, snapshot_2024_targets = _build_2024_snapshot_summary(adjustment_snapshot_2024_path)
    program_catalog_summary, program_catalog_targets = _build_2026_catalog_summary(program_catalog_2026_path)
    stats_2025_full_summary, stats_2025_full_targets = _build_2025_full_stats_summary(adjustment_stats_2025_full_path)
    announcement_2025_summary, announcement_2025_targets = _build_2025_announcement_summary(adjustment_announcement_2025_path)
    landing_2024_summary, landing_2024_targets = _build_2024_landing_summary(adjustment_landing_2024_path)

    deduped_targets = _dedupe_targets(
        [
            *snapshot_2024_targets,
            *program_catalog_targets,
            *stats_2025_full_targets,
            *announcement_2025_targets,
            *landing_2024_targets,
        ]
    )

    outputs = {
        snapshot_2024_output: snapshot_2024_summary,
        program_catalog_output: program_catalog_summary,
        stats_2025_full_output: stats_2025_full_summary,
        announcement_2025_output: announcement_2025_summary,
        landing_2024_output: landing_2024_summary,
    }
    for output_path, payload in outputs.items():
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    targets_output.parent.mkdir(parents=True, exist_ok=True)
    targets_output.write_text(json.dumps(deduped_targets, ensure_ascii=False, indent=2), encoding="utf-8")

    return {
        "snapshot_2024_rows": snapshot_2024_summary["total_rows"],
        "program_catalog_2026_rows": program_catalog_summary["total_rows"],
        "stats_2025_full_rows": stats_2025_full_summary["total_rows"],
        "announcement_2025_rows": announcement_2025_summary["total_rows"],
        "landing_2024_rows": landing_2024_summary["total_rows"],
        "target_rows": len(deduped_targets),
        "target_unique_schools": len({str(row["school_name"]).strip() for row in deduped_targets}),
        "targets_output": str(targets_output),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Build extended Stage A summaries and school targets from additional workbooks.")
    parser.add_argument("--adjustment-snapshot-2024", required=True)
    parser.add_argument("--program-catalog-2026", required=True)
    parser.add_argument("--adjustment-stats-2025-full", required=True)
    parser.add_argument("--adjustment-announcement-2025", required=True)
    parser.add_argument("--adjustment-landing-2024", required=True)
    parser.add_argument("--snapshot-2024-output", default="docs/data/adjustment_snapshot_2024_0414_summary.json")
    parser.add_argument("--program-catalog-output", default="docs/data/admission_program_catalog_2026_summary.json")
    parser.add_argument("--stats-2025-full-output", default="docs/data/adjustment_stats_2025_full_summary.json")
    parser.add_argument("--announcement-2025-output", default="docs/data/adjustment_announcement_2025_summary.json")
    parser.add_argument("--landing-2024-output", default="docs/data/adjustment_landing_2024_summary.json")
    parser.add_argument("--targets-output", default="docs/data/adjustment_expanded_priority_targets_2024_2026.json")
    args = parser.parse_args()

    result = build_outputs(
        adjustment_snapshot_2024_path=Path(args.adjustment_snapshot_2024).expanduser().resolve(),
        program_catalog_2026_path=Path(args.program_catalog_2026).expanduser().resolve(),
        adjustment_stats_2025_full_path=Path(args.adjustment_stats_2025_full).expanduser().resolve(),
        adjustment_announcement_2025_path=Path(args.adjustment_announcement_2025).expanduser().resolve(),
        adjustment_landing_2024_path=Path(args.adjustment_landing_2024).expanduser().resolve(),
        snapshot_2024_output=Path(args.snapshot_2024_output).resolve(),
        program_catalog_output=Path(args.program_catalog_output).resolve(),
        stats_2025_full_output=Path(args.stats_2025_full_output).resolve(),
        announcement_2025_output=Path(args.announcement_2025_output).resolve(),
        landing_2024_output=Path(args.landing_2024_output).resolve(),
        targets_output=Path(args.targets_output).resolve(),
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
