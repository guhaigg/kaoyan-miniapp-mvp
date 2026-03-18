from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


TAG_RE = re.compile(r"<[^>]+>")


def _clean_review_preview(text: str) -> str:
    cleaned = text.replace("&lt;", "<").replace("&gt;", ">")
    cleaned = cleaned.replace("<br>", " ").replace("<br/>", " ").replace("<br />", " ")
    cleaned = TAG_RE.sub(" ", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return cleaned[:160]


def build_summary(*, input_path: Path, output_path: Path, top_schools: int = 20) -> dict:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    next(rows, None)

    school_counts: Counter[str] = Counter()
    duplicate_keys: Counter[tuple[str, str]] = Counter()
    htmlish_rows = 0
    url_rows = 0
    bad_rows = 0
    total_rows = 0
    samples: list[dict] = []

    for row in rows:
        school = str(row[0] or "").strip()
        college = str(row[1] or "").strip()
        name = str(row[2] or "").strip()
        review = str(row[3] or "").strip()
        total_rows += 1
        school_counts[school] += 1
        duplicate_keys[(school, name)] += 1
        if school in {"", "-"} or name in {"", "-"}:
            bad_rows += 1
        if "<" in review or "&lt;" in review or "<br" in review.lower():
            htmlish_rows += 1
        if "http://" in review or "https://" in review:
            url_rows += 1
        if len(samples) < 5 and school not in {"", "-"} and name not in {"", "-"} and review:
            samples.append(
                {
                    "school_name": school,
                    "college_name": college if college not in {"", "-"} else None,
                    "mentor_name": name,
                    "review_preview": _clean_review_preview(review),
                }
            )

    summary = {
        "source_file": str(input_path),
        "sheet_name": worksheet.title,
        "total_rows": total_rows,
        "unique_schools": len(school_counts),
        "htmlish_rows": htmlish_rows,
        "url_rows": url_rows,
        "bad_rows": bad_rows,
        "duplicate_name_keys": sum(1 for _, count in duplicate_keys.items() if count > 1),
        "top_schools": [
            {
                "rank": rank,
                "school_name": school_name,
                "school_category": "导师公开评价",
                "adjustment_count": count,
                "top_departments": [],
            }
            for rank, (school_name, count) in enumerate(school_counts.most_common(top_schools), start=1)
        ],
        "samples": samples,
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description="Build mentor review summary JSON from workbook.")
    parser.add_argument("--input", required=True, help="Path to the mentor review workbook")
    parser.add_argument(
        "--output",
        default="docs/data/mentor_review_summary.json",
        help="Path to write the mentor review summary JSON",
    )
    parser.add_argument("--top-schools", type=int, default=20)
    args = parser.parse_args()
    summary = build_summary(
        input_path=Path(args.input).expanduser().resolve(),
        output_path=Path(args.output).resolve(),
        top_schools=args.top_schools,
    )
    print(
        json.dumps(
            {
                "total_rows": summary["total_rows"],
                "unique_schools": summary["unique_schools"],
                "htmlish_rows": summary["htmlish_rows"],
                "output": str(Path(args.output).resolve()),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
