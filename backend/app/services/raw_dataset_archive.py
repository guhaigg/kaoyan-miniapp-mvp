from __future__ import annotations

import base64
import gzip
import hashlib
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from ..models import RawDatasetArchive


def _coerce_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _extract_xlsx_metadata(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_names = list(workbook.sheetnames)
    primary_sheet_name = sheet_names[0] if sheet_names else None
    total_rows = None
    total_columns = None
    header_row: list[str] = []
    preview_rows: list[list[str]] = []

    if primary_sheet_name:
        worksheet = workbook[primary_sheet_name]
        total_rows = worksheet.max_row
        total_columns = worksheet.max_column
        rows = worksheet.iter_rows(values_only=True)
        first_row = next(rows, None)
        if first_row is not None:
            header_row = [_coerce_cell(value) for value in first_row]
        for _ in range(3):
            row = next(rows, None)
            if row is None:
                break
            preview_rows.append([_coerce_cell(value) for value in row])

    return {
        "workbook_format": "xlsx",
        "sheet_names": sheet_names,
        "primary_sheet_name": primary_sheet_name,
        "total_rows": total_rows,
        "total_columns": total_columns,
        "header_row": header_row,
        "preview_rows": preview_rows,
    }


def _extract_xls_metadata(path: Path) -> dict[str, Any]:
    try:
        import xlrd  # type: ignore
    except ImportError as exc:  # pragma: no cover - validated in integration use
        raise RuntimeError("xlrd is required to archive .xls datasets") from exc

    workbook = xlrd.open_workbook(path)
    sheet_names = workbook.sheet_names()
    primary_sheet_name = sheet_names[0] if sheet_names else None
    total_rows = None
    total_columns = None
    header_row: list[str] = []
    preview_rows: list[list[str]] = []

    if primary_sheet_name:
        sheet = workbook.sheet_by_name(primary_sheet_name)
        total_rows = sheet.nrows
        total_columns = sheet.ncols
        if sheet.nrows > 0:
            header_row = [_coerce_cell(sheet.cell_value(0, col)) for col in range(sheet.ncols)]
        for row_idx in range(1, min(sheet.nrows, 4)):
            preview_rows.append([_coerce_cell(sheet.cell_value(row_idx, col)) for col in range(sheet.ncols)])

    return {
        "workbook_format": "xls",
        "sheet_names": sheet_names,
        "primary_sheet_name": primary_sheet_name,
        "total_rows": total_rows,
        "total_columns": total_columns,
        "header_row": header_row,
        "preview_rows": preview_rows,
    }


def extract_workbook_metadata(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _extract_xlsx_metadata(path)
    if suffix == ".xls":
        return _extract_xls_metadata(path)
    raise ValueError(f"unsupported workbook format: {suffix or 'unknown'}")


def archive_dataset_file(
    db: Session,
    *,
    dataset_key: str,
    title: str,
    dataset_type: str,
    input_path: Path,
    summary_json: dict[str, Any] | None = None,
    notes: str | None = None,
) -> RawDatasetArchive:
    path = input_path.expanduser().resolve()
    raw_bytes = path.read_bytes()
    payload = base64.b64encode(gzip.compress(raw_bytes)).decode("ascii")
    metadata = extract_workbook_metadata(path)
    sha256 = hashlib.sha256(raw_bytes).hexdigest()

    archive = db.query(RawDatasetArchive).filter(RawDatasetArchive.dataset_key == dataset_key).one_or_none()
    if archive is None:
        archive = RawDatasetArchive(dataset_key=dataset_key)
        db.add(archive)

    archive.title = title
    archive.dataset_type = dataset_type
    archive.source_filename = path.name
    archive.source_path = str(path)
    archive.workbook_format = metadata["workbook_format"]
    archive.file_sha256 = sha256
    archive.file_size_bytes = len(raw_bytes)
    archive.storage_encoding = "gzip_base64"
    archive.raw_file_payload = payload
    archive.sheet_names = list(metadata["sheet_names"])
    archive.primary_sheet_name = metadata["primary_sheet_name"]
    archive.total_rows = metadata["total_rows"]
    archive.total_columns = metadata["total_columns"]
    archive.header_row = list(metadata["header_row"])
    archive.preview_rows = list(metadata["preview_rows"])
    archive.summary_json = dict(summary_json or {})
    archive.notes = notes
    db.flush()
    return archive
